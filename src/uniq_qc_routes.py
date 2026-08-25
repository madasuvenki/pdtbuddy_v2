"""
UniqQC Routes — Unique CR Quality Dashboard
Integrates the UniQ SQL Model dashboard into PDT Buddy.

Backend source is PDT Buddy MySQL:
- target metadata from pdt_stats_dashboard.dashboard_status via dashboard_common
- CR quality data from per-BU {db_prefix}_overallcrs / {db_prefix}_overall_crs tables
"""

from __future__ import annotations

import datetime as _dt
import csv
import io
import logging
import re
import zipfile
from typing import Any

from flask import Blueprint, jsonify, render_template, request, send_file, session

logger = logging.getLogger(__name__)

uniq_qc_bp = Blueprint("uniq_qc_bp", __name__)

_IDENT_RE = re.compile(r"^[A-Za-z0-9_]+$")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _get_mysql_conn():
    """Get MySQL connection using PDT Buddy's existing utility."""
    from src.utils import get_mysql_connection_db

    return get_mysql_connection_db()


def _safe_ident(value: str) -> str:
    """Return a safely quoted MySQL identifier."""
    value = str(value or "").strip()
    if not value or not _IDENT_RE.match(value):
        raise ValueError(f"Unsafe SQL identifier: {value!r}")
    return f"`{value}`"


def _fq(schema: str, table: str) -> str:
    return f"{_safe_ident(schema)}.{_safe_ident(table)}"


def _norm_key(value: Any) -> str:
    return str(value or "").strip().lower()


def _target_prefix_candidates(target_name: str, info: dict | None = None) -> list[str]:
    """Return possible DB table base prefixes for a target.

    PDT Buddy stores the authoritative prefix in dashboard_status.db_name, exposed
    as db_name/db_prefix by dashboard_common. The target key is only a fallback.
    """
    info = info or {}
    raw_values = [
        info.get("overall_crs_table"),
        info.get("overallcrs_table"),
        info.get("db_prefix"),
        info.get("db_name"),
        info.get("target_name"),
        target_name,
        info.get("display_name"),
        info.get("target_display"),
        info.get("sp_name"),
    ]

    out: list[str] = []
    seen: set[str] = set()
    for raw in raw_values:
        val = str(raw or "").strip()
        if not val:
            continue

        # If a fully-qualified or explicit table name was configured, keep only table part.
        if "." in val:
            val = val.split(".")[-1]
        val = val.strip("`")

        base = val.lower().replace("-", "_").replace(" ", "_").replace(".", "_")
        for suffix in ("_overallcrs", "_overall_crs"):
            if base.endswith(suffix):
                base = base[: -len(suffix)]
                break

        base = re.sub(r"[^a-z0-9_]+", "_", base).strip("_")
        if base and base not in seen:
            seen.add(base)
            out.append(base)

    return out


def _table_exists(conn, schema: str, table_name: str) -> bool:
    """Check if a table exists in the given schema."""
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT COUNT(*)
            FROM information_schema.tables
            WHERE table_schema=%s AND table_name=%s
            """,
            (schema, table_name),
        )
        row = cur.fetchone()
        return (row[0] if row else 0) > 0
    except Exception:
        return False
    finally:
        try:
            cur.close()
        except Exception:
            pass


def _table_columns(conn, schema: str, table_name: str) -> dict[str, str]:
    """Return lower-case column name -> actual column name."""
    cur = conn.cursor(dictionary=True)
    try:
        cur.execute(f"SHOW COLUMNS FROM {_fq(schema, table_name)}")
        return {str(r["Field"]).lower(): str(r["Field"]) for r in (cur.fetchall() or [])}
    finally:
        cur.close()


def _pick_col(cols: dict[str, str], *exact_or_contains: str) -> str | None:
    for token in exact_or_contains:
        token_l = token.lower()
        if token_l in cols:
            return cols[token_l]
    for token in exact_or_contains:
        token_l = token.lower()
        for low, actual in cols.items():
            if token_l in low:
                return actual
    return None


def _get_targets_config() -> dict[str, dict]:
    from dashboard_common import get_targets_config, load_metadata_config

    targets = get_targets_config() or {}
    if targets:
        return targets
    # Fallback when globals were not warmed yet.
    return (load_metadata_config(active_only=True).get("TARGETS_CONFIG") or {})


def _get_business_units() -> dict[str, dict]:
    from dashboard_common import get_business_units, load_metadata_config

    bus = get_business_units() or {}
    if bus:
        return bus
    return (load_metadata_config(active_only=True).get("BUSINESS_UNITS") or {})


def _get_schema_for_target(target_name: str) -> str:
    from dashboard_common import get_schema_for_target

    return get_schema_for_target(target_name) or ""


def _get_bu_for_target(target_name: str, info: dict | None = None) -> str:
    from dashboard_common import get_bu_for_target

    try:
        return get_bu_for_target(target_name) or str((info or {}).get("bu") or "Unknown")
    except Exception:
        return str((info or {}).get("bu") or "Unknown")


def _display_name(target_name: str, info: dict | None = None) -> str:
    info = info or {}
    return str(
        info.get("sp_name")
        or info.get("display_name")
        or info.get("target_display")
        or target_name
    )


def _find_overallcrs_table(conn, schema: str, target_name: str, info: dict | None = None) -> str | None:
    """Find the actual overall CR table for a target."""
    for base in _target_prefix_candidates(target_name, info):
        for table in (f"{base}_overallcrs", f"{base}_overall_crs"):
            if _table_exists(conn, schema, table):
                return table
    return None


def _reported_team_expr(cols: dict[str, str]) -> str:
    team_col = _pick_col(cols, "reported_team", "test_team", "team")
    if team_col:
        return f"COALESCE({_safe_ident(team_col)}, '')"
    return "''"


def _crid_col(cols: dict[str, str]) -> str | None:
    return _pick_col(cols, "crid", "mapped_cr", "mapped_crs", "cr")


def _date_col(cols: dict[str, str]) -> str | None:
    return _pick_col(cols, "date", "created_date", "submitted_date", "updated_date")


def _fetch_statistics_for_target(
    conn,
    schema: str,
    target_name: str,
    info: dict | None = None,
) -> dict | None:
    """Query the target overallcrs table and compute UniQ dashboard statistics."""
    table = _find_overallcrs_table(conn, schema, target_name, info)
    if not table:
        return None

    fq_table = _fq(schema, table)

    try:
        cols = _table_columns(conn, schema, table)
        cr_col = _crid_col(cols)
        team_expr = _reported_team_expr(cols)

        cur = conn.cursor(dictionary=True)

        if cr_col:
            count_expr = f"COUNT(DISTINCT NULLIF(TRIM({_safe_ident(cr_col)}), ''))"
        else:
            count_expr = "COUNT(*)"

        cur.execute(f"SELECT {count_expr} AS cnt FROM {fq_table}")
        total_crs = int((cur.fetchone() or {}).get("cnt") or 0)
        if total_crs == 0:
            return None

        cur.execute(
            f"""
            SELECT {team_expr} AS reported_team, {count_expr} AS cnt
            FROM {fq_table}
            GROUP BY {team_expr}
            """
        )

        team_counts: dict[str, int] = {}
        for row in cur.fetchall() or []:
            key = str(row.get("reported_team") or "").strip()
            team_counts[key] = int(row.get("cnt") or 0)

        pdt_unique = team_counts.get("PDT_Unique", 0)
        pdt_overlap = team_counts.get("PDT_Reported", 0)
        pdt_reported = pdt_unique + pdt_overlap
        other_teams = max(total_crs - pdt_reported, 0)

        bu = _get_bu_for_target(target_name, info)
        return {
            "target": target_name,
            "sp_name": _display_name(target_name, info),
            "total_crs": total_crs,
            "pdt_reported": pdt_reported,
            "pdt_unique": pdt_unique,
            "pdt_unique_pct": round((pdt_unique / total_crs * 100), 2) if total_crs else 0.0,
            "overlapping": pdt_overlap,
            "other_teams": other_teams,
            "seen_only_by_other_team": other_teams,
            "group": bu,
            "bu": bu,
            "table": table,
            "schema": schema,
            "is_total": False,
        }
    except Exception as e:
        logger.warning("[UniqQC] Error fetching stats for %s.%s target=%s: %s", schema, table, target_name, e)
        return None
    finally:
        try:
            cur.close()
        except Exception:
            pass


def _fetch_subsystem_for_target(
    conn,
    schema: str,
    target_name: str,
    mode: str,
    info: dict | None = None,
) -> dict[str, list]:
    """Fetch subsystem breakdown in frontend format: {subsystems: [], counts: []}."""
    table = _find_overallcrs_table(conn, schema, target_name, info)
    if not table:
        return {"subsystems": [], "counts": []}

    try:
        cols = _table_columns(conn, schema, table)
        subs_col = _pick_col(cols, "subsystem", "subs", "sub_system")
        if not subs_col:
            return {"subsystems": [], "counts": []}

        cr_col = _crid_col(cols)
        count_expr = f"COUNT(DISTINCT NULLIF(TRIM({_safe_ident(cr_col)}), ''))" if cr_col else "COUNT(*)"

        where_parts = [f"{_safe_ident(subs_col)} IS NOT NULL", f"TRIM({_safe_ident(subs_col)}) <> ''"]
        if mode == "unique":
            where_parts.append("reported_team = 'PDT_Unique'")
        elif mode == "reported":
            where_parts.append("reported_team IN ('PDT_Reported', 'PDT_Unique')")

        # Only apply reported_team predicates when the column exists.
        if ("reported_team" not in cols) and any("reported_team" in p for p in where_parts):
            where_parts = [p for p in where_parts if "reported_team" not in p]

        cur = conn.cursor(dictionary=True)
        cur.execute(
            f"""
            SELECT {_safe_ident(subs_col)} AS subs, {count_expr} AS cnt
            FROM {_fq(schema, table)}
            WHERE {" AND ".join(where_parts)}
            GROUP BY {_safe_ident(subs_col)}
            ORDER BY cnt DESC
            LIMIT 100
            """
        )

        rows = cur.fetchall() or []
        return {
            "subsystems": [str(r.get("subs") or "Unknown") for r in rows],
            "counts": [int(r.get("cnt") or 0) for r in rows],
        }
    except Exception as e:
        logger.warning("[UniqQC] Error fetching subsystem data for %s: %s", target_name, e)
        return {"subsystems": [], "counts": []}
    finally:
        try:
            cur.close()
        except Exception:
            pass


def _fetch_all_subsystems(mode: str) -> dict[str, dict[str, list]]:
    conn = None
    try:
        conn = _get_mysql_conn()
        targets_config = _get_targets_config()
        result: dict[str, dict[str, list]] = {}

        for target_name, info in targets_config.items():
            schema = _get_schema_for_target(target_name)
            if not schema:
                continue
            data = _fetch_subsystem_for_target(conn, schema, target_name, mode, info)
            if data.get("subsystems"):
                result[target_name] = data
                display = _display_name(target_name, info)
                if display != target_name:
                    result[display] = data
        return result
    finally:
        try:
            if conn:
                conn.close()
        except Exception:
            pass


def _fetch_cr_details_for_target(conn, schema: str, target_name: str, cr_type: str, info: dict | None = None) -> list:
    """Fetch individual CR details for the modal."""
    table = _find_overallcrs_table(conn, schema, target_name, info)
    if not table:
        return []

    try:
        cols = _table_columns(conn, schema, table)

        crid_col = _crid_col(cols)
        area_col = _pick_col(cols, "area")
        subs_col = _pick_col(cols, "subsystem", "subs", "sub_system")
        func_col = _pick_col(cols, "func", "functionality")
        status_col = _pick_col(cols, "status")
        team_col = _pick_col(cols, "reported_team", "test_team", "team")

        if not crid_col:
            return []

        if team_col:
            if cr_type == "other_teams":
                where = f"{_safe_ident(team_col)} NOT IN ('PDT_Reported', 'PDT_Unique')"
            elif cr_type == "overlapping":
                where = f"{_safe_ident(team_col)} = 'PDT_Reported'"
            elif cr_type == "pdt_unique":
                where = f"{_safe_ident(team_col)} = 'PDT_Unique'"
            else:
                where = "1=1"
        else:
            where = "1=1"

        select_cols = [f"{_safe_ident(crid_col)} AS crid"]
        if area_col:
            select_cols.append(f"{_safe_ident(area_col)} AS area")
        if subs_col:
            select_cols.append(f"{_safe_ident(subs_col)} AS subsystem")
        if func_col:
            select_cols.append(f"{_safe_ident(func_col)} AS func")
        if status_col:
            select_cols.append(f"{_safe_ident(status_col)} AS status")

        cur = conn.cursor(dictionary=True)
        cur.execute(
            f"""
            SELECT {", ".join(select_cols)}
            FROM {_fq(schema, table)}
            WHERE {where}
              AND {_safe_ident(crid_col)} IS NOT NULL
              AND TRIM({_safe_ident(crid_col)}) <> ''
            LIMIT 1000
            """
        )

        return [
            {
                "crid": str(r.get("crid") or ""),
                "area": r.get("area", ""),
                "subs": r.get("subsystem", ""),
                "func": r.get("func", ""),
                "status": r.get("status", ""),
                "pdt_tag": "",
            }
            for r in (cur.fetchall() or [])
        ]
    except Exception as e:
        logger.warning("[UniqQC] Error fetching CR details for %s: %s", target_name, e)
        return []
    finally:
        try:
            cur.close()
        except Exception:
            pass


def _build_statistics_payload() -> list[dict]:
    conn = None
    try:
        conn = _get_mysql_conn()
        targets_config = _get_targets_config()

        all_stats: list[dict] = []
        bu_totals: dict[str, dict[str, int]] = {}

        for target_name, info in targets_config.items():
            schema = _get_schema_for_target(target_name)
            if not schema:
                continue

            stat = _fetch_statistics_for_target(conn, schema, target_name, info)
            if not stat:
                continue

            all_stats.append(stat)
            group = stat["group"]
            totals = bu_totals.setdefault(
                group,
                {
                    "total_crs": 0,
                    "pdt_reported": 0,
                    "pdt_unique": 0,
                    "overlapping": 0,
                    "other_teams": 0,
                    "seen_only_by_other_team": 0,
                },
            )
            for key in totals:
                totals[key] += int(stat.get(key) or 0)

        for group, totals in bu_totals.items():
            total_crs = totals["total_crs"]
            pdt_unique = totals["pdt_unique"]
            all_stats.append(
                {
                    "target": group,
                    "sp_name": f"{group} (Total)",
                    "total_crs": total_crs,
                    "pdt_reported": totals["pdt_reported"],
                    "pdt_unique": pdt_unique,
                    "pdt_unique_pct": round(pdt_unique / total_crs * 100, 2) if total_crs else 0.0,
                    "overlapping": totals["overlapping"],
                    "other_teams": totals["other_teams"],
                    "seen_only_by_other_team": totals["seen_only_by_other_team"],
                    "group": group,
                    "bu": group,
                    "is_total": True,
                }
            )

        return all_stats
    finally:
        try:
            if conn:
                conn.close()
        except Exception:
            pass


def _latest_data_date() -> dict[str, Any]:
    """Return latest unique_cr/dashboard update date from dashboard_status metadata."""
    targets = _get_targets_config()
    latest = None
    bu_dates: dict[str, str] = {}

    for target, info in targets.items():
        raw = (info or {}).get("unique_cr_last_update") or (info or {}).get("dashboard_latest_update")
        if not raw:
            continue
        if isinstance(raw, (_dt.datetime, _dt.date)):
            dt = raw if isinstance(raw, _dt.datetime) else _dt.datetime.combine(raw, _dt.time.min)
        else:
            try:
                dt = _dt.datetime.fromisoformat(str(raw).replace("Z", "+00:00")).replace(tzinfo=None)
            except Exception:
                continue
        if latest is None or dt > latest:
            latest = dt
        bu = _get_bu_for_target(target, info)
        prev = bu_dates.get(bu)
        if not prev or str(dt) > prev:
            bu_dates[bu] = dt.strftime("%Y-%m-%d")

    if latest:
        return {
            "last_data_date": latest.strftime("%Y-%m-%d"),
            "formatted_date": latest.strftime("%b %d, %Y"),
            "bu_data_dates": bu_dates,
        }

    today = _dt.datetime.now()
    return {
        "last_data_date": today.strftime("%Y-%m-%d"),
        "formatted_date": today.strftime("%b %d, %Y"),
        "bu_data_dates": bu_dates,
    }


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------

@uniq_qc_bp.route("/uniq_qc")
def uniq_qc_dashboard():
    """Render the UniqQC dashboard page."""
    return render_template("uniq_qc_dashboard.html")


@uniq_qc_bp.route("/api/uniq_qc/data")
def api_uniq_qc_data():
    """Standalone UniQ-compatible data endpoint."""
    try:
        return jsonify(_build_statistics_payload())
    except Exception as e:
        logger.exception("[UniqQC] Error in data API")
        return jsonify({"error": str(e)}), 500


@uniq_qc_bp.route("/api/uniq_qc/statistics")
def api_uniq_qc_statistics():
    """Structured statistics endpoint for PDT Buddy consumers."""
    try:
        return jsonify(
            {
                "success": True,
                "statistics": _build_statistics_payload(),
                **_latest_data_date(),
            }
        )
    except Exception as e:
        logger.exception("[UniqQC] Error in statistics API")
        return jsonify({"success": False, "error": str(e)}), 500


@uniq_qc_bp.route("/api/uniq_qc/subsystems")
def api_uniq_qc_subsystems():
    """Return PDT Unique subsystem data for all targets, or one target when requested."""
    target = request.args.get("target", "").strip()
    if not target:
        return jsonify(_fetch_all_subsystems("unique"))

    try:
        conn = _get_mysql_conn()
        info = _get_targets_config().get(target) or {}
        schema = _get_schema_for_target(target)
        data = _fetch_subsystem_for_target(conn, schema, target, "unique", info) if schema else {"subsystems": [], "counts": []}
        return jsonify(
            {
                "success": True,
                "subsystem_data": dict(zip(data["subsystems"], data["counts"])),
                "subsystems": data["subsystems"],
                "counts": data["counts"],
            }
        )
    except Exception as e:
        logger.exception("[UniqQC] Error in subsystems API")
        return jsonify({"error": str(e)}), 500
    finally:
        try:
            conn.close()
        except Exception:
            pass


@uniq_qc_bp.route("/api/uniq_qc/subsystems-reported")
def api_uniq_qc_subsystems_reported():
    """Return PDT Reported + PDT Unique subsystem data for all targets."""
    try:
        return jsonify(_fetch_all_subsystems("reported"))
    except Exception as e:
        logger.exception("[UniqQC] Error in subsystems-reported API")
        return jsonify({"error": str(e)}), 500


@uniq_qc_bp.route("/api/uniq_qc/subsystems-overall")
def api_uniq_qc_subsystems_overall():
    """Return overall subsystem data for all targets."""
    try:
        return jsonify(_fetch_all_subsystems("overall"))
    except Exception as e:
        logger.exception("[UniqQC] Error in subsystems-overall API")
        return jsonify({"error": str(e)}), 500


@uniq_qc_bp.route("/api/uniq_qc/hierarchy")
@uniq_qc_bp.route("/api/uniq_qc/bu-hierarchy")
def api_uniq_qc_hierarchy():
    """Return BU hierarchy for the dashboard filter."""
    try:
        business_units = _get_business_units()
        targets_config = _get_targets_config()

        hierarchy: dict[str, dict[str, list]] = {}
        for bu_key, bu_info in (business_units or {}).items():
            targets = list((bu_info or {}).get("targets") or [])
            if not targets:
                targets = [t for t, info in targets_config.items() if str((info or {}).get("bu") or "").upper() == str(bu_key).upper()]
            if targets:
                hierarchy[str(bu_key)] = {str(t): [] for t in targets}

        # Template's older JS expects the raw object from /bu-hierarchy.
        if request.path.endswith("/bu-hierarchy"):
            return jsonify(hierarchy)

        return jsonify({"success": True, "bu_hierarchy": hierarchy})
    except Exception as e:
        logger.exception("[UniqQC] Error in hierarchy API")
        return jsonify({"error": str(e)}), 500


@uniq_qc_bp.route("/api/uniq_qc/last-data-date")
def api_uniq_qc_last_data_date():
    return jsonify(_latest_data_date())


@uniq_qc_bp.route("/api/uniq_qc/refresh-data", methods=["POST"])
def api_uniq_qc_refresh_data():
    """No file refresh is needed; data is read live from PDT Buddy DB tables."""
    return jsonify({"success": True, "message": "UniqQC reads live PDT Buddy database tables."})


@uniq_qc_bp.route("/api/uniq_qc/cr_details")
def api_uniq_qc_cr_details_query():
    target = request.args.get("target", "")
    cr_type = request.args.get("cr_type", "other_teams")
    return _cr_details_response(target, cr_type)


@uniq_qc_bp.route("/api/cr-details/<path:target>/<cr_type>")
@uniq_qc_bp.route("/api/uniq_qc/cr-details/<path:target>/<cr_type>")
def api_uniq_qc_cr_details_path(target: str, cr_type: str):
    return _cr_details_response(target, cr_type)


def _cr_details_response(target: str, cr_type: str):
    """Return individual CR details for the modal."""
    if not target:
        return jsonify({"crs": [], "count": 0})

    conn = None
    try:
        conn = _get_mysql_conn()
        targets = _get_targets_config()
        info = targets.get(target) or {}

        # Resolve display-name clicks back to canonical target key.
        if not info:
            for key, cfg in targets.items():
                if _display_name(key, cfg).lower() == target.lower():
                    target = key
                    info = cfg
                    break

        schema = _get_schema_for_target(target)
        if not schema:
            return jsonify({"crs": [], "count": 0})

        crs = _fetch_cr_details_for_target(conn, schema, target, cr_type, info)
        return jsonify({"crs": crs, "count": len(crs)})
    except Exception as e:
        logger.exception("[UniqQC] Error in cr_details API")
        return jsonify({"error": str(e)}), 500
    finally:
        try:
            if conn:
                conn.close()
        except Exception:
            pass


def _csv_bytes_for_target(target: str) -> tuple[bytes, str] | None:
    """Return CSV bytes for one target's overallcrs table."""
    conn = None
    try:
        conn = _get_mysql_conn()
        targets = _get_targets_config()
        info = targets.get(target) or {}

        if not info:
            for key, cfg in targets.items():
                if _display_name(key, cfg).lower() == target.lower():
                    target = key
                    info = cfg
                    break

        schema = _get_schema_for_target(target)
        if not schema:
            return None

        table = _find_overallcrs_table(conn, schema, target, info)
        if not table:
            return None

        cur = conn.cursor(dictionary=True)
        cur.execute(f"SELECT * FROM {_fq(schema, table)} LIMIT 200000")
        rows = cur.fetchall() or []
        headers = list(rows[0].keys()) if rows else list(_table_columns(conn, schema, table).values())

        text_buf = io.StringIO(newline="")
        writer = csv.DictWriter(text_buf, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)

        filename = f"{_display_name(target, info).replace(' ', '_')}_overallcrs.csv"
        return text_buf.getvalue().encode("utf-8-sig"), filename
    finally:
        try:
            cur.close()
        except Exception:
            pass
        try:
            if conn:
                conn.close()
        except Exception:
            pass


@uniq_qc_bp.route("/api/uniq_qc/download_excel")
def api_uniq_qc_download_excel_zip():
    """Download selected BU overallcrs data as a ZIP of CSV files.

    The imported standalone UI calls this "Excel"; CSV keeps the endpoint
    dependency-free and is directly openable in Excel.
    """
    selected_bu = str(request.args.get("bu") or "").strip()
    targets = _get_targets_config()

    buf = io.BytesIO()
    added = 0
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for target, info in targets.items():
            bu = _get_bu_for_target(target, info)
            if selected_bu and selected_bu.lower() not in {bu.lower(), "automotive" if bu.upper() == "AUTO" else bu.lower()}:
                continue
            csv_payload = _csv_bytes_for_target(target)
            if not csv_payload:
                continue
            data, filename = csv_payload
            zf.writestr(filename, data)
            added += 1

    if not added:
        return jsonify({"error": "No overallcrs data available for selected BU"}), 404

    buf.seek(0)
    safe_bu = re.sub(r"[^A-Za-z0-9_]+", "_", selected_bu or "UniqQC").strip("_")
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"{safe_bu}_UniqQC_overallcrs_csv.zip",
        mimetype="application/zip",
    )


@uniq_qc_bp.route("/api/uniq_qc/download_excel/<path:target>")
def api_uniq_qc_download_excel_target(target: str):
    csv_payload = _csv_bytes_for_target(target)
    if not csv_payload:
        return jsonify({"error": "No overallcrs data available for selected target"}), 404

    data, filename = csv_payload
    return send_file(
        io.BytesIO(data),
        as_attachment=True,
        download_name=filename,
        mimetype="text/csv",
    )


@uniq_qc_bp.route("/api/uniq_qc/cr-summary", methods=["POST"])
def api_uniq_qc_cr_summary():
    """Lightweight CR summary endpoint for the RCA button.

    Uses existing Orbit/QGenie plumbing when available; otherwise returns a safe
    structured fallback so the modal does not break.
    """
    payload = request.get_json(silent=True) or {}
    cr_number = str(payload.get("message") or payload.get("crid") or "").strip()
    if not cr_number:
        return jsonify({"content": "No CR number provided."}), 400

    cr_number_clean = cr_number.replace("CR", "").replace("cr", "").strip()
    try:
        from orbit_client import fetch_cr

        cr = fetch_cr(cr_number_clean, use_cache=True) or {}
        title = cr.get("title") or cr.get("headline") or cr.get("subject") or ""
        status = cr.get("status") or cr.get("state") or ""
        priority = cr.get("priority") or ""
        content = (
            f"**CR{cr_number_clean}**\n"
            f"Status: {status or 'N/A'}\n"
            f"Priority: {priority or 'N/A'}\n"
            f"Summary: {title or 'Orbit details unavailable.'}"
        )
        return jsonify({"content": content})
    except Exception as exc:
        logger.info("[UniqQC] CR summary fallback for CR%s: %s", cr_number_clean, exc)
        return jsonify(
            {
                "content": (
                    f"**CR{cr_number_clean}**\n"
                    "RCA summary is not available from Orbit/QGenie in this session. "
                    "Open the Orbit link for full CR details."
                )
            }
        )


@uniq_qc_bp.route("/api/uniq_qc/download_ppt")
def api_uniq_qc_download_ppt():
    """Generate and return a PPT file with the Venn diagram statistics."""
    try:
        all_stats = [s for s in _build_statistics_payload() if not s.get("is_total")]
        if not all_stats:
            return jsonify({"error": "No data available"}), 404

        try:
            from pptx import Presentation
            from pptx.dml.color import RGBColor
            from pptx.enum.text import PP_ALIGN
            from pptx.util import Inches, Pt

            prs = Presentation()
            prs.slide_width = Inches(13.33)
            prs.slide_height = Inches(7.5)
            slide_layout = prs.slide_layouts[6]

            slide = prs.slides.add_slide(slide_layout)
            tx_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(1.5))
            p = tx_box.text_frame.paragraphs[0]
            p.text = "UniqQC — Unique CR Quality Dashboard"
            p.font.size = Pt(32)
            p.font.bold = True
            p.font.color.rgb = RGBColor(0x0A, 0x0E, 0x27)
            p.alignment = PP_ALIGN.CENTER

            tx_box2 = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(11), Inches(0.8))
            p2 = tx_box2.text_frame.paragraphs[0]
            p2.text = f"Generated: {_dt.datetime.now().strftime('%Y-%m-%d %H:%M')}"
            p2.font.size = Pt(14)
            p2.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            p2.alignment = PP_ALIGN.CENTER

            slide2 = prs.slides.add_slide(slide_layout)
            title_box = slide2.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.7))
            p3 = title_box.text_frame.paragraphs[0]
            p3.text = "CR Statistics Summary"
            p3.font.size = Pt(24)
            p3.font.bold = True
            p3.font.color.rgb = RGBColor(0x0A, 0x0E, 0x27)

            y_pos = Inches(1.2)
            for stat in all_stats[:20]:
                row_text = (
                    f"{stat['sp_name']:30s}  "
                    f"Total: {stat['total_crs']:5d}  "
                    f"PDT Reported: {stat['pdt_reported']:5d}  "
                    f"PDT Unique: {stat['pdt_unique']:5d}  "
                    f"({stat['pdt_unique_pct']:.1f}%)"
                )
                row_box = slide2.shapes.add_textbox(Inches(0.5), y_pos, Inches(12), Inches(0.35))
                row_p = row_box.text_frame.paragraphs[0]
                row_p.text = row_text
                row_p.font.size = Pt(10)
                row_p.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
                y_pos += Inches(0.38)
                if y_pos > Inches(7.0):
                    break

            buf = io.BytesIO()
            prs.save(buf)
            buf.seek(0)

            return send_file(
                buf,
                as_attachment=True,
                download_name="UniqQC_Dashboard.pptx",
                mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )

        except ImportError:
            return jsonify({"error": "python-pptx not installed. Install with: pip install python-pptx"}), 500

    except Exception as e:
        logger.exception("[UniqQC] Error generating PPT")
        return jsonify({"error": str(e)}), 500