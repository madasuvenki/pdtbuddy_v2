"""
Automotive Top 5 / Top 10 CRs feature (JIRA QIPLPDT-11137).

Provides:
  * GET  /auto/top_crs                       -> dedicated page
  * GET  /api/auto/top_crs                    -> ranked CR JSON (most-recent first)
  * GET  /api/auto/top_crs/targets            -> Automotive target dropdown options
  * GET  /auto/top_crs/download               -> PowerPoint export of the ranked CRs

Top CRs are selected by recency: CRs are ordered by their most recent activity
timestamp (``jira_date__last_instance``) descending. Users can pick Top 5 or
Top 10 and filter by crash type (System / SSR / Process crashes).

Data source: the family-level ``<target>_unique_crs`` table is the single source
of truth for the CR list and all display columns (per the PDT lead's guidance --
it is the complete, correctly-ordered universe of unique CRs). Crash type is
classified from the ``cr_title`` text (which carries the ``_SystemCrash`` /
``ProcessDump``/``[LINUX_PROCESS_CRASH]`` / ``SSR`` signatures) and, when the CR
also exists in ``<target>_overallcrs``, is reinforced by that table's ``label``
tag. We no longer rank off ``overallcrs`` because it is a crash-instance subset
that omits some of the most recent CRs.
"""

from __future__ import annotations

import io
import logging
import re
from datetime import date, datetime

from flask import (
    Blueprint,
    jsonify,
    render_template,
    request,
    send_file,
)
from flask_login import login_required

import dashboard_common as dc
from src.utils import get_mysql_connection_db

logger = logging.getLogger(__name__)

auto_top_crs_bp = Blueprint("auto_top_crs", __name__)

# ---------------------------------------------------------------------------
# Crash-type classification
# ---------------------------------------------------------------------------
# The AUTO overallcrs ``label`` column carries an authoritative crash-type tag
# of the form ``<build>_SystemCrash`` / ``_SSR`` / ``_ProcessCrash`` / ``_others``
# (e.g. ``SA8255P.HGY.4.1.8.0_SystemCrash``). That explicit tag is the primary
# signal. Only when no tag is present do we fall back to keyword heuristics that
# mirror the rest of PDT Buddy (open_jiras_section.html / live_status_view_sp.html).

# Authoritative label suffix tags -> crash type.
_CRASH_TAG_RE = re.compile(
    r"_(SystemCrash|SSRCrash|SSR|ProcessCrash|ProcessDump|Process|others|Others)\b",
    re.IGNORECASE,
)
_TAG_TO_TYPE = {
    "systemcrash": "system",
    "ssr": "ssr",
    "ssrcrash": "ssr",
    "process": "process",
    "processcrash": "process",
    "processdump": "process",
    "others": "other",
}

# Keyword fallbacks (used only when no explicit tag is found).
_SSR_KEYWORDS = (
    "ssr", "subsystem restart", "subsystem_restart", "sub-system restart",
)
_PROCESS_KEYWORDS = (
    "process crash", "process_crash", "processcrash", "processdump",
    "tombstone", "app not responding", "anr", "fatal exception",
    "java.lang.", "linux_process_crash", "userspace",
)
_SYSTEM_KEYWORDS = (
    "systemcrash", "system crash", "system_crash", "ramdump", "ram dump",
    "err_fatal", "err fatal",
)


def _classify_crash_type(*texts: str) -> str:
    """Return one of 'system', 'ssr', 'process' or 'other' from CR text/labels.

    Resolution order:
      1. Explicit ``_SystemCrash`` / ``_SSR`` / ``_ProcessCrash`` / ``_others``
         label suffix tags (authoritative). If a CR carries multiple tags across
         its rows, the most severe wins: system > ssr > process > other.
      2. Keyword heuristics as a fallback when no tag is present.
    """
    blob = " ".join(str(t or "") for t in texts)

    # 1) Authoritative label tags.
    found = set()
    for m in _CRASH_TAG_RE.finditer(blob):
        mapped = _TAG_TO_TYPE.get(m.group(1).lower())
        if mapped:
            found.add(mapped)
    if found:
        for sev in ("system", "ssr", "process", "other"):
            if sev in found:
                return sev

    # 2) Keyword fallback.
    low = blob.lower()
    if any(k in low for k in _SSR_KEYWORDS):
        return "ssr"
    if any(k in low for k in _PROCESS_KEYWORDS):
        return "process"
    if any(k in low for k in _SYSTEM_KEYWORDS):
        return "system"
    return "other"


# ---------------------------------------------------------------------------
# Automotive target + sub-target discovery
# ---------------------------------------------------------------------------

# A sub-target keyword appearing as a ``_adas_`` / ``_ivi_`` / ``_flex_`` segment.
_SUBTARGET_RE = re.compile(r"(?:^|_)(adas|ivi|flex)(?:_|$)")
# Optional target variant segment (e.g. monaco_HGY_adas...) before the sub-target.
_TABLE_PARSE_RE = re.compile(r"^(?:(hgy|hqx|la)_)?(adas|ivi|flex)(?:_(.+))?$")


def _all_unique_crs_tables() -> set[str]:
    """Return the set of all ``*_unique_crs`` table names in pdt_stats_auto."""
    conn = get_mysql_connection_db()
    if not conn:
        return set()
    try:
        cur = conn.cursor()
        cur.execute("SHOW TABLES FROM pdt_stats_auto LIKE '%_unique_crs'")
        tables = {str(t) for (t,) in cur.fetchall()}
        cur.close()
        return tables
    except Exception:
        logger.exception("Failed to list unique_crs tables")
        return set()
    finally:
        try:
            conn.close()
        except Exception:
            pass


def _family_prefixes(tables: set[str]) -> list[str]:
    """Return the family prefixes (the actual selectable targets).

    A family is a prefix that (a) has its own base ``<prefix>_unique_crs`` table
    and (b) contains no adas/ivi/flex sub-target token. e.g. ``nord_hgy``,
    ``lemans_hqx``, ``lemans_la``, ``monaco``, ``nord_hqx``.

    Note Monaco: the base table is ``monaco_unique_crs`` while its sub-target
    tables are named ``monaco_hgy_*`` / ``monaco_hqx_*``. Because ``monaco_hgy``
    has NO base table, it is not a family -- those sub-tables attach to the
    ``monaco`` family with the hgy/hqx variant folded into the sub-target key.
    """
    prefixes = {t[: -len("_unique_crs")] for t in tables}
    return sorted(
        p for p in prefixes
        if not _SUBTARGET_RE.search(p)  # no adas/ivi/flex token
    )


def _parse_unique_table(table: str, families: list[str]) -> tuple[str | None, str, str | None]:
    """Parse a ``*_unique_crs`` table into (family, subtarget, version).

    - family: the owning selectable target (e.g. ``nord_hgy``, ``monaco``).
    - subtarget: 'base' for the family table itself; otherwise the sub-target
      key, which includes an optional variant segment folded in for Monaco-style
      naming (e.g. ``hqx_adas``, ``hgy_adas``) or plain ``adas``/``ivi``/``flex``.
    - version: None when the sub-target table has no version segment.
    Returns (None, '', None) if the table does not belong to any known family.
    """
    rest = table[: -len("_unique_crs")]
    # Longest matching family prefix wins (so 'nord_hgy' beats 'nord').
    fam = None
    for f in sorted(families, key=len, reverse=True):
        if rest == f or rest.startswith(f + "_"):
            fam = f
            break
    if fam is None:
        return (None, "", None)
    rem = rest[len(fam):].lstrip("_")
    if not rem:
        return (fam, "base", None)
    m = _TABLE_PARSE_RE.match(rem)
    if not m:
        return (fam, "", None)
    variant, sub, ver = m.group(1), m.group(2), m.group(3)
    # Fold the variant (hgy/hqx/la) into the sub-target key so Monaco's
    # monaco_hqx_adas / monaco_hgy_adas become distinct sub-targets of 'monaco'.
    sub_key = f"{variant}_{sub}" if variant else sub
    return (fam, sub_key, ver)


def _build_selection_tree() -> dict:
    """Return {family: {"display": str, "subtargets": {sub: [versions...]}}}.

    ``versions`` is a sorted list; a version-less sub-target table contributes
    the sentinel ``""`` (empty string) to that list.
    """
    tables = _all_unique_crs_tables()
    families = _family_prefixes(tables)
    tree: dict = {}
    for fam in families:
        tree[fam] = {"subtargets": {}}
    for t in tables:
        fam, sub, ver = _parse_unique_table(t, families)
        if not fam or fam not in tree:
            continue
        if sub in ("", "base"):
            continue
        tree[fam]["subtargets"].setdefault(sub, [])
        tree[fam]["subtargets"][sub].append(ver or "")
    for fam in tree:
        for sub in tree[fam]["subtargets"]:
            tree[fam]["subtargets"][sub] = sorted(set(tree[fam]["subtargets"][sub]))
    return tree


def _auto_overallcrs_targets() -> list[dict]:
    """Return Automotive targets that have a family ``*_unique_crs`` table.

    Each entry: {"key": <family_prefix>, "display_name": <label>}.
    Families are discovered from the base ``<prefix>_unique_crs`` tables.
    """
    tables = _all_unique_crs_tables()
    if not tables:
        return []
    families = _family_prefixes(tables)
    base_prefixes = {t[: -len("_unique_crs")] for t in tables}
    # Only families that actually own a base table are selectable targets.
    qualified = [f for f in families if f in base_prefixes]

    # Map prefixes to friendly display names from TARGETS_CONFIG when available.
    try:
        dc.update_global_targets_config()
        cfg = dc.get_targets_config() or {}
    except Exception:
        cfg = {}
    prefix_to_display: dict[str, str] = {}
    for _tkey, _info in cfg.items():
        if str((_info or {}).get("bu", "")).upper() != "AUTO":
            continue
        _prefix = str((_info or {}).get("db_prefix", _tkey) or _tkey).lower()
        _disp = str((_info or {}).get("display_name") or _tkey)
        prefix_to_display.setdefault(_prefix, _disp)

    def _pretty(prefix: str) -> str:
        if prefix in prefix_to_display:
            return prefix_to_display[prefix]
        return prefix.replace("_", " ").upper()

    return [{"key": p, "display_name": _pretty(p)} for p in sorted(qualified)]


def _target_subtargets(target_key: str) -> dict:
    """Return the sub-target/version options for one family target.

    Sub-target keys may include a variant prefix for Monaco-style naming
    (e.g. ``hqx_adas``); the label is prettified accordingly (e.g. 'HQX ADAS').
    """
    tree = _build_selection_tree()
    fam = tree.get(target_key, {})
    subs = fam.get("subtargets", {})
    out = [{"key": "all", "label": "All (Base)", "versions": []}]
    for sub in sorted(subs):
        versions = [v for v in subs[sub] if v]  # drop the version-less sentinel
        out.append(
            {
                "key": sub,
                "label": sub.replace("_", " ").upper(),
                "versions": versions,
                "has_versionless": "" in subs[sub],
            }
        )
    return {"subtargets": out}


def _valid_target_key(target_key: str) -> str | None:
    """Return the target_key only if it maps to a real unique_crs table."""
    key = re.sub(r"[^a-z0-9_]", "", str(target_key or "").strip().lower())
    if not key:
        return None
    valid = {t["key"] for t in _auto_overallcrs_targets()}
    return key if key in valid else None


# ---------------------------------------------------------------------------
# Top CR data
# ---------------------------------------------------------------------------

def _date_text(value) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    raw = str(value or "").strip()
    m = re.match(r"^(\d{4}-\d{2}-\d{2})", raw)
    return m.group(1) if m else raw


def _clean(value) -> str:
    """Return a trimmed string, treating None/NA/nan placeholders as blank."""
    s = str(value if value is not None else "").strip()
    if s.lower() in ("", "none", "nan", "na", "null"):
        return ""
    return s


def _unique_crs_table(target_key: str) -> str | None:
    """Return the family-level ``<target>_unique_crs`` table name if it exists."""
    conn = get_mysql_connection_db()
    if not conn:
        return None
    try:
        cur = conn.cursor()
        cur.execute(
            "SHOW TABLES FROM pdt_stats_auto LIKE %s",
            (f"{target_key}_unique_crs",),
        )
        row = cur.fetchone()
        cur.close()
        return f"{target_key}_unique_crs" if row else None
    except Exception:
        logger.exception("Failed to look up unique_crs table for %s", target_key)
        return None
    finally:
        try:
            conn.close()
        except Exception:
            pass


def _overallcrs_label_map(target_key: str) -> dict:
    """Return {CR_ID_UPPER: concatenated label text} from ``<target>_overallcrs``.

    Used only to reinforce crash-type classification for CRs that also appear in
    the overallcrs table (which carries the authoritative ``_SystemCrash`` /
    ``_SSR`` / ``_ProcessCrash`` label tags). CRs absent here are classified
    from their ``cr_title`` alone.
    """
    conn = get_mysql_connection_db()
    if not conn:
        return {}
    out: dict[str, str] = {}
    try:
        cur = conn.cursor(dictionary=True)
        cur.execute(
            f"""
            SELECT crid, GROUP_CONCAT(COALESCE(label, '') SEPARATOR ';') AS labels
            FROM pdt_stats_auto.`{target_key}_overallcrs`
            WHERE crid IS NOT NULL AND crid <> ''
            GROUP BY crid
            """
        )
        for r in cur.fetchall() or []:
            key = _clean(r.get("crid")).upper()
            if key:
                out[key] = str(r.get("labels") or "")
        cur.close()
    except Exception:
        # overallcrs may not exist for every target -- that's fine.
        logger.debug("No overallcrs label map for %s", target_key, exc_info=True)
    finally:
        try:
            conn.close()
        except Exception:
            pass
    return out


def _resolve_source_tables(target_key: str, subtarget: str, version: str) -> list[str]:
    """Return the list of ``*_unique_crs`` tables to read for a selection.

    - subtarget 'all'            -> [<target>_unique_crs]            (base family)
    - subtarget X, version 'all' -> every table of that sub-target (all versions,
                                     including the version-less one)
    - subtarget X, version V      -> the single matching table
    The sub-target key may include a Monaco-style variant prefix (e.g.
    ``hqx_adas``). Only tables that actually exist are returned.
    """
    tables = _all_unique_crs_tables()
    subtarget = (subtarget or "all").lower()
    version = (version or "all").lower()

    if subtarget == "all":
        base = f"{target_key}_unique_crs"
        return [base] if base in tables else []

    families = _family_prefixes(tables)
    out = []
    for t in tables:
        fam, sub, ver = _parse_unique_table(t, families)
        if fam != target_key or sub != subtarget:
            continue
        if version and version != "all" and (ver or "") != version:
            continue
        out.append(t)
    return sorted(out)


def _subtarget_membership(target_key: str) -> dict:
    """Return {CR_ID_UPPER: set(subtarget_label)} for the sub-target tables.

    Used to derive the 'Common' column: a CR's membership is the set of
    sub-targets whose ``*_unique_crs`` tables contain that CR id (option (b):
    derived from real table membership, not ``seen_in_targets``). The label is
    the sub-target key prettified (e.g. ``adas`` -> 'ADAS', ``hqx_adas`` ->
    'HQX ADAS').
    """
    tables = _all_unique_crs_tables()
    families = _family_prefixes(tables)
    membership: dict[str, set] = {}
    conn = get_mysql_connection_db()
    if not conn:
        return membership
    try:
        cur = conn.cursor()
        for t in tables:
            fam, sub, _ver = _parse_unique_table(t, families)
            if fam != target_key or sub in ("", "base"):
                continue
            label = sub.replace("_", " ").upper()
            try:
                cur.execute(
                    f"SELECT DISTINCT cr FROM pdt_stats_auto.`{t}` "
                    f"WHERE cr IS NOT NULL AND cr <> ''"
                )
                for (cr,) in cur.fetchall():
                    key = _clean(cr).upper()
                    if key:
                        membership.setdefault(key, set()).add(label)
            except Exception:
                logger.debug("Failed reading membership from %s", t, exc_info=True)
        cur.close()
    except Exception:
        logger.exception("Failed to build sub-target membership for %s", target_key)
    finally:
        try:
            conn.close()
        except Exception:
            pass
    return membership


def _format_common(subs: set, all_subs: set) -> str:
    """Format the Common column: 'All' if present in every sub-target of the
    family, else the comma-joined sub-target names (e.g. 'ADAS, IVI'). Empty if
    the CR is in no sub-target table."""
    if not subs:
        return ""
    order = ["ADAS", "IVI", "FLEX"]
    present = [s for s in order if s in subs] + sorted(subs - set(order))
    if all_subs and subs >= all_subs and len(all_subs) >= 2:
        return "All"
    return ", ".join(present)


# A build is a Release Build (RB) when its metabuild string carries an
# ``R<digit>`` segment (e.g. ``...5.1.7.0.R1-00145-...``); otherwise it is a
# Mainline (ML) build (``...C1-...``, ``...-01227-...`` etc.).
_RB_BUILD_RE = re.compile(r"[._-]R\d", re.IGNORECASE)


def _rb_ml_map(target_key: str) -> dict:
    """Return {CR_ID_UPPER: 'RB' | 'ML' | 'RB/ML'} from the family ``*_jiras``
    table's ``metabuild`` column.

    A CR seen only on R* builds -> 'RB'; only on non-R* builds -> 'ML'; on both
    -> 'RB/ML'. CRs with no jira rows are absent from the map (rendered blank).
    """
    conn = get_mysql_connection_db()
    if not conn:
        return {}
    out: dict[str, str] = {}
    try:
        cur = conn.cursor()
        cur.execute(
            f"""
            SELECT cr, metabuild
            FROM pdt_stats_auto.`{target_key}_jiras`
            WHERE cr IS NOT NULL AND cr <> ''
            """
        )
        has_rb: dict[str, bool] = {}
        has_ml: dict[str, bool] = {}
        for cr, metabuild in cur.fetchall():
            key = _clean(cr).upper()
            if not key:
                continue
            if _RB_BUILD_RE.search(str(metabuild or "")):
                has_rb[key] = True
            else:
                has_ml[key] = True
        cur.close()
        for key in set(has_rb) | set(has_ml):
            rb, ml = has_rb.get(key), has_ml.get(key)
            out[key] = "RB/ML" if (rb and ml) else ("RB" if rb else "ML")
    except Exception:
        # jiras table may not exist for every target -- that's fine (blank).
        logger.debug("No RB/ML map for %s", target_key, exc_info=True)
    finally:
        try:
            conn.close()
        except Exception:
            pass
    return out


def _pdt_priority(occurrence) -> str:
    """Derive PDT priority from Jira count: >10 -> P1, else P2.

    Non-numeric occurrence (e.g. 'Dup') is treated as 0 -> P2.
    """
    return "P1" if _occurrence_sort_key(occurrence) > 10 else "P2"


def _occurrence_sort_key(value) -> int:
    """Numeric sort helper for ``cr_occurrence`` ('Dup'/'NA'/'' -> 0)."""
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return 0


def _fetch_top_crs(
    target_key: str,
    crash_types: set,
    limit: int,
    subtarget: str = "all",
    version: str = "all",
) -> list[dict]:
    """Return the Top-N CRs for a target (optionally a sub-target/version).

    List source is the resolved ``*_unique_crs`` table(s), ordered by
    ``jira_date__last_instance`` descending (matches the DB view). Crash type is
    classified from ``cr_title`` + ``overallcrs.label``. The 'Common' column is
    derived from which ADAS/IVI/FLEX tables contain the CR. PDT priority is
    derived from Jira count (>10 -> P1 else P2).

    crash_types: subset of {'system', 'ssr', 'process'}. Empty set = no filter.
    """
    source_tables = _resolve_source_tables(target_key, subtarget, version)
    if not source_tables:
        return []
    conn = get_mysql_connection_db()
    if not conn:
        return []
    raw: list[dict] = []
    try:
        cur = conn.cursor(dictionary=True)
        union_sql = " UNION ALL ".join(
            f"""SELECT cr, cr_occurrence, pdt_priority_tag, cr_age, cr_title,
                       cr_area, cr_subsystem, cr_functionality, cr_date, cr_status,
                       cr_notes, jira_date__last_instance, image, cr_category,
                       parent_cr
                FROM pdt_stats_auto.`{t}`
                WHERE cr IS NOT NULL AND cr <> ''"""
            for t in source_tables
        )
        cur.execute(
            f"""
            SELECT * FROM ( {union_sql} ) AS u
            ORDER BY (jira_date__last_instance IS NULL),
                     jira_date__last_instance DESC
            """
        )
        raw = cur.fetchall() or []
        cur.close()
    except Exception:
        logger.exception("Failed to query unique_crs for %s/%s/%s", target_key, subtarget, version)
        return []
    finally:
        try:
            conn.close()
        except Exception:
            pass

    label_map = _overallcrs_label_map(target_key)
    membership = _subtarget_membership(target_key)
    rb_ml_map = _rb_ml_map(target_key)
    all_subs = set()
    for _subs in membership.values():
        all_subs |= _subs

    rows_out: list[dict] = []
    seen: set = set()
    for r in raw:
        cr_id = _clean(r.get("cr"))
        if not cr_id or cr_id.upper() in seen:
            continue
        title = _clean(r.get("cr_title"))
        crash = _classify_crash_type(
            title,
            label_map.get(cr_id.upper(), ""),
            _clean(r.get("cr_functionality")),
            _clean(r.get("cr_subsystem")),
        )
        if crash_types and crash not in crash_types:
            continue
        seen.add(cr_id.upper())
        occurrence = _clean(r.get("cr_occurrence"))
        parent = _clean(r.get("parent_cr"))
        # A duplicate CR points to its parent; don't show self-references.
        is_dup = occurrence.lower() == "dup" or (parent and parent.upper() != cr_id.upper())
        parent_cr = parent if (parent and parent.upper() != cr_id.upper()) else ""
        rows_out.append(
            {
                "cr": cr_id,
                "parent_cr": parent_cr,
                "is_dup": bool(parent_cr),
                "occurrence": occurrence,
                "priority": _pdt_priority(occurrence),
                "crash_type": crash,
                "age": _clean(r.get("cr_age")),
                "title": title,
                "area": _clean(r.get("cr_area")),
                "subsystem": _clean(r.get("cr_subsystem")),
                "functionality": _clean(r.get("cr_functionality")),
                "meta_id": _clean(r.get("image")),
                "rb_ml": rb_ml_map.get(cr_id.upper(), ""),
                "common": _format_common(membership.get(cr_id.upper(), set()), all_subs),
                "cr_date": _date_text(r.get("cr_date")),
                "last_seen": _date_text(r.get("jira_date__last_instance")),
                "status": _clean(r.get("cr_status")),
                "notes": _clean(r.get("cr_notes")),
                "scenario": "",  # user-editable; not sourced from DB
            }
        )
        if len(rows_out) >= limit:
            break
    return rows_out


def _parse_request_params():
    """Extract (target_key, crash_types, limit, subtarget, version) from args."""
    target_key = _valid_target_key(request.args.get("target", ""))
    try:
        top_n = int(request.args.get("top", "5"))
    except (TypeError, ValueError):
        top_n = 5
    limit = 10 if top_n >= 10 else 5

    requested = {
        c.strip().lower()
        for c in (request.args.get("crash_types", "") or "").split(",")
        if c.strip()
    }
    crash_types = requested & {"system", "ssr", "process"}

    # Sub-target key may include a Monaco-style variant prefix (e.g. 'hqx_adas').
    subtarget = re.sub(r"[^a-z_]", "", request.args.get("subtarget", "all").strip().lower()) or "all"
    version = re.sub(r"[^a-z0-9_]", "", request.args.get("version", "all").strip().lower()) or "all"
    return target_key, crash_types, limit, subtarget, version


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------

@auto_top_crs_bp.route("/auto/top_crs")
@login_required
def auto_top_crs_page():
    targets = _auto_overallcrs_targets()
    return render_template(
        "auto_top_crs.html",
        bu_name="Automotive",
        targets=targets,
    )


@auto_top_crs_bp.route("/api/auto/top_crs/targets")
@login_required
def auto_top_crs_targets():
    return jsonify({"targets": _auto_overallcrs_targets()})


@auto_top_crs_bp.route("/api/auto/top_crs/subtargets")
@login_required
def auto_top_crs_subtargets():
    """Return sub-target + version options for a given target."""
    target_key = _valid_target_key(request.args.get("target", ""))
    if not target_key:
        return jsonify({"success": False, "error": "Invalid or missing target."}), 400
    data = _target_subtargets(target_key)
    return jsonify({"success": True, "target": target_key, **data})


@auto_top_crs_bp.route("/api/auto/top_crs")
@login_required
def auto_top_crs_data():
    target_key, crash_types, limit, subtarget, version = _parse_request_params()
    if not target_key:
        return jsonify({"success": False, "error": "Invalid or missing target."}), 400
    rows = _fetch_top_crs(target_key, crash_types, limit, subtarget, version)
    return jsonify(
        {
            "success": True,
            "target": target_key,
            "subtarget": subtarget,
            "version": version,
            "top": limit,
            "crash_types": sorted(crash_types) or ["system", "ssr", "process"],
            "count": len(rows),
            "rows": rows,
        }
    )


@auto_top_crs_bp.route("/auto/top_crs/download", methods=["GET", "POST"])
@login_required
def auto_top_crs_download():
    target_key, crash_types, limit, subtarget, version = _parse_request_params()
    if not target_key:
        return jsonify({"success": False, "error": "Invalid or missing target."}), 400
    rows = _fetch_top_crs(target_key, crash_types, limit, subtarget, version)

    # User-edited Scenario text is posted as a JSON map {CR_ID: scenario}.
    scenarios: dict = {}
    if request.method == "POST":
        payload = request.get_json(silent=True) or {}
        raw_map = payload.get("scenarios") if isinstance(payload, dict) else None
        if isinstance(raw_map, dict):
            scenarios = {str(k).upper(): str(v or "") for k, v in raw_map.items()}
    if scenarios:
        for r in rows:
            override = scenarios.get(str(r.get("cr", "")).upper())
            if override is not None:
                r["scenario"] = override

    display_name = next(
        (t["display_name"] for t in _auto_overallcrs_targets() if t["key"] == target_key),
        target_key.upper(),
    )
    # Append the sub-target / version to the display name when not the base.
    if subtarget and subtarget != "all":
        display_name += f" — {subtarget.replace('_', ' ').upper()}"
        if version and version != "all":
            display_name += f" ({version})"
    crash_label = (
        ", ".join(c.upper() for c in sorted(crash_types))
        if crash_types
        else "All Crash Types"
    )
    buf = _build_top_crs_ppt(display_name, rows, limit, crash_label)
    _slug = re.sub(r"[^A-Za-z0-9]+", "_", f"{target_key}_{subtarget}_{version}").strip("_")
    filename = f"Top_{limit}_CRs_{_slug}_{datetime.now():%Y%m%d}.pptx"
    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )


# ---------------------------------------------------------------------------
# PowerPoint export
# ---------------------------------------------------------------------------

def _build_top_crs_ppt(target_display: str, rows: list[dict], limit: int, crash_label: str) -> io.BytesIO:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    BLUE = RGBColor(0x1F, 0x5F, 0x91)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    BLACK = RGBColor(0x00, 0x00, 0x00)
    ROW = RGBColor(0xE9, 0xED, 0xF3)
    ROW_ALT = RGBColor(0xF4, 0xEC, 0xF2)

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(12.7), Inches(0.6))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f"{target_display} — PDT Top {limit} CRs ({crash_label})"
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = BLUE

    sub_box = slide.shapes.add_textbox(Inches(0.3), Inches(0.78), Inches(12.7), Inches(0.3))
    sp = sub_box.text_frame.paragraphs[0]
    srun = sp.add_run()
    srun.text = f"Selected by most recent activity  |  Generated {datetime.now():%Y-%m-%d %H:%M}"
    srun.font.size = Pt(10)
    srun.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    headers = [
        "S.No", "CR ID", "Jira count\n(Cumulative)", "PDT\nPriority",
        "Crash Type", "CR age", "CR Title", "CR Area",
        "Last seen ML and\nRB META ID", "RB/ML", "Common",
        "CR date", "Last Seen", "CR Status", "Debug Notes", "Scenario",
    ]
    col_widths = [0.4, 0.9, 0.7, 0.55, 0.7, 0.5, 2.4, 1.0, 1.0, 0.5, 0.5, 0.7, 0.7, 0.7, 1.9, 1.1]

    n_rows = max(1, len(rows)) + 1
    n_cols = len(headers)
    table_shape = slide.shapes.add_table(
        n_rows, n_cols, Inches(0.12), Inches(1.15), Inches(13.1), Inches(6.0)
    )
    tbl = table_shape.table
    total_w = sum(col_widths)
    for ci, cw in enumerate(col_widths):
        tbl.columns[ci].width = int(Inches(13.1) * (cw / total_w))

    def _set(cell, text, size=6.5, bold=False, color=BLACK, fill=None, align=PP_ALIGN.CENTER):
        if fill is not None:
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
        ctf = cell.text_frame
        ctf.clear()
        ctf.word_wrap = True
        ctf.margin_left = Pt(1.5)
        ctf.margin_right = Pt(1.5)
        ctf.margin_top = Pt(0.5)
        ctf.margin_bottom = Pt(0.5)
        para = ctf.paragraphs[0]
        para.alignment = align
        r = para.add_run()
        r.text = str(text or "")
        r.font.name = "Arial"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color

    for ci, h in enumerate(headers):
        _set(tbl.cell(0, ci), h, size=6.5, bold=True, color=WHITE, fill=BLUE)

    # Columns that read better left-aligned (free text).
    _left_cols = {6, 14, 15}

    if rows:
        for ri, row in enumerate(rows, start=1):
            bg = ROW if ri % 2 else ROW_ALT
            # Duplicates show their parent CR in slash form (child / parent).
            cr_disp = row.get("cr", "")
            if row.get("parent_cr"):
                cr_disp = f"{cr_disp} / {row.get('parent_cr')}"
            # Jira count: duplicates note the parent they roll up into.
            occ_disp = str(row.get("occurrence", "") or "")
            if row.get("is_dup") and row.get("parent_cr"):
                occ_disp = f"Dup of {row.get('parent_cr')}"
            vals = [
                str(ri),
                cr_disp,
                occ_disp,
                row.get("priority", "") or "-",
                row.get("crash_type", "").upper(),
                row.get("age", "") or "-",
                row.get("title", ""),
                row.get("area", ""),
                row.get("meta_id", "") or "-",
                row.get("rb_ml", "") or "-",
                row.get("common", "") or "-",
                row.get("cr_date", ""),
                row.get("last_seen", ""),
                row.get("status", ""),
                row.get("notes", "") or "No Debug Notes",
                row.get("scenario", "") or "-",
            ]
            for ci, val in enumerate(vals):
                align = PP_ALIGN.LEFT if ci in _left_cols else PP_ALIGN.CENTER
                _set(tbl.cell(ri, ci), val, size=6.0, fill=bg, align=align)
    else:
        _set(tbl.cell(1, 0), "No CRs found for this selection.", size=8, fill=ROW, align=PP_ALIGN.LEFT)
        for ci in range(1, n_cols):
            _set(tbl.cell(1, ci), "", size=8, fill=ROW)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf
