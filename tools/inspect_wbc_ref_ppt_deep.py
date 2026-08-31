from pptx import Presentation

p = r"\\sphere\pdtstats\CR_TAT_Script\PDTBuddy\PPT\PDT_Amboseli.LE.2.0_WBC_Core_25082026.pptx"
prs = Presentation(p)
print("slides", len(prs.slides), "size", prs.slide_width, prs.slide_height)
for i, s in enumerate(prs.slides, 1):
    print(f"\n=== SLIDE {i} ===")
    for j, sh in enumerate(s.shapes, 1):
        typ = getattr(sh, "shape_type", "")
        name = getattr(sh, "name", "")
        print(f"{j:02d} type={typ} name={name} x={sh.left} y={sh.top} w={sh.width} h={sh.height}")
        if hasattr(sh, "text") and sh.text.strip():
            print("   TEXT:", sh.text.strip().replace("\n", " | ")[:800])
        if getattr(sh, "has_table", False):
            tbl = sh.table
            print(f"   TABLE rows={len(tbl.rows)} cols={len(tbl.columns)}")
            for r_idx, r in enumerate(list(tbl.rows)[:8], 1):
                print("    ", r_idx, " || ".join(c.text.strip().replace("\n", " ") for c in list(r.cells)[:10])[:1000])
