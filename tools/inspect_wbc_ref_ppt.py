from pptx import Presentation

p = r"\\sphere\pdtstats\CR_TAT_Script\PDTBuddy\PPT\PDT_Amboseli.LE.2.0_WBC_Core_25082026.pptx"
prs = Presentation(p)
print("slides", len(prs.slides), "size", prs.slide_width, prs.slide_height)
for i, s in enumerate(prs.slides, 1):
    print("\n--- SLIDE", i, "---")
    texts = []
    for sh in s.shapes:
        if hasattr(sh, "text") and sh.text.strip():
            texts.append(sh.text.strip().replace("\n", " | "))
        if getattr(sh, "has_table", False):
            rows = []
            for r in list(sh.table.rows)[:5]:
                rows.append(" || ".join(c.text.strip().replace("\n", " ") for c in list(r.cells)[:8]))
            texts.append("TABLE: " + " / ".join(rows))
    print("\n".join(texts[:30]))